import json
import openpyxl
import os
import shutil
from datetime import datetime
from typing import Dict
import logging
from app.config import Config

logger = logging.getLogger(__name__)

def normalize_name(name: str) -> str:
    """Нормализует имя для сравнения."""
    return name.strip().lower().replace('\n', '').replace('\r', '')

def normalize_code(code: str) -> str:
    """Нормализует табельный номер."""
    return str(code).strip().lower()

def check_with_excel(workers_json_path: str, excel_path: str = None):
    """
    Проверяет данные workers.json по Excel файлу.
    Создает два отчета о расхождениях.
    """
    if excel_path is None:
        excel_path = Config.EXCEL_FILE

    try:
        # Загрузка Excel через openpyxl
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws = wb.active

        # Извлекаем имена (столбец 3) и табельные номера (столбец 2)
        excel_names = []
        excel_ids = []
        for row in range(2, ws.max_row + 1):
            name_cell = ws.cell(row=row, column=3).value
            id_cell = ws.cell(row=row, column=2).value
            if not name_cell:
                break
            excel_names.append(normalize_name(str(name_cell)))
            if id_cell:
                excel_ids.append(normalize_code(str(id_cell)))

        excel_names_set = set(excel_names)
        excel_ids_set = set(excel_ids)

        with open(workers_json_path, 'r', encoding='utf-8') as f:
            workers_data = json.load(f)

        # Проверка по именам
        missing_by_name = {
            emp_id: data
            for emp_id, data in workers_data.items()
            if normalize_name(data.get('name', '')) not in excel_names_set
        }

        # Проверка по табельным номерам
        missing_by_code = {
            emp_id: data
            for emp_id, data in workers_data.items()
            if normalize_code(data.get('matchcode', '')) not in excel_ids_set
        }

        # Сохраняем отчеты (пути из конфига)
        save_missing_report(missing_by_name, Config.MISSING_FROM_EXCEL, "по ФИО")
        save_missing_report(missing_by_code, Config.MISSING_BY_MATCHCODE, "по табельным номерам")

    except FileNotFoundError:
        logger.warning(f"Файл {excel_path} не найден. Проверка пропущена.")
    except Exception as e:
        logger.error(f"Ошибка при проверке данных: {e}")
        raise

def save_missing_report(data: Dict, filename: str, report_type: str):
    """Сохраняет отчет о недостающих данных."""
    if data:
        with open(filename, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        logger.info(f"Сохранено {len(data)} сотрудников, которых нет в Excel {report_type} в {filename}")
    else:
        logger.info(f"Все сотрудники присутствуют в Excel {report_type}")

def archive_workers_json(source_file: str = None):
    """Архивирует JSON файл."""
    if source_file is None:
        source_file = Config.OUTPUT_JSON
    if not os.path.exists(source_file):
        logger.warning(f"Файл {source_file} не найден для архивации")
        return
    os.makedirs(Config.ARCHIVE_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    archive_file = os.path.join(Config.ARCHIVE_DIR, f"workers_{timestamp}.json")
    shutil.copy2(source_file, archive_file)
    logger.info(f"Файл сохранен в архив: {archive_file}")

def archive_workers_csv(csv_file_path: str = None):
    """Архивирует CSV файл."""
    if csv_file_path is None:
        csv_file_path = Config.OUTPUT_CSV
    if not os.path.exists(csv_file_path):
        logger.warning(f"Файл {csv_file_path} не найден для архивации")
        return
    os.makedirs(Config.ARCHIVE_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    archive_filename = f"workers_{timestamp}.csv"
    archive_path = os.path.join(Config.ARCHIVE_DIR, archive_filename)
    shutil.copy2(csv_file_path, archive_path)
    logger.info(f"Файл {csv_file_path} заархивирован как {archive_path}")

def archive_files():
    """Архивирует все основные файлы."""
    archive_workers_json()
    archive_workers_csv()